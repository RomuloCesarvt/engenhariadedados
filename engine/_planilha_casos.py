"""Gera a planilha de vendas SDR do período (uso: python engine/_planilha_casos.py 2026-09-01 2026-09-30)."""
import json
import os
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from engine import casos_sdr  # noqa: E402

inicio, fim = sys.argv[1], sys.argv[2]
arq = os.path.join("output", "analises", "periodos", f"{inicio.replace('-', '')}_{fim.replace('-', '')}.json")
analise = json.load(open(arq, encoding="utf-8"))["analise"]
todas = casos_sdr.todas(analise)


def reais(v):
    return "R$ " + f"{v:,.0f}".replace(",", ".")


CORES = {"remanescente": "D6DBDF", "sdr": "D5F5E3", "perdida": "FADBD8", "sdr_cv_sem_evidencia": "FFE2B3", "parcial": "FFF8DC", "sem_rastro": "FFFFFF"}
VERDE = Font(bold=True, color="FFFFFF")
FUNDO = PatternFill("solid", fgColor="1B4F38")
mes = f"{inicio[8:]}/{inicio[5:7]} a {fim[8:]}/{fim[5:7]}/{fim[:4]}"

wb = Workbook()
ws = wb.active
ws.title = "Resumo"
ws.append([f"Vendas SDR — {mes}"])
ws.append([])
ws.append(["Vendas do período", len(todas), reais(sum(t["valor"] for t in todas))])
for k, rotulo in casos_sdr.VEREDITOS.items():
    grupo = [t for t in todas if t["veredito"] == k]
    ws.append([rotulo, len(grupo), reais(sum(t["valor"] for t in grupo))])
ws.append([])
for linha in (
    "Critério da gestão: venda SDR é quando o SDR (ou a assistente virtual) atende o lead até o agendamento ou o repasse para um corretor,",
    "o lead vem de origem online (paga ou orgânica, a lista do BI) e isso é recente (último atendimento até 90 dias antes da venda), numa praça com ação online no ar.",
    "Praça remanescente (sem investimento online nos 90 dias antes da venda) fica fora da análise; está na última aba só para conferência.",
    "Se o corretor registrou o cliente antes do SDR/assistente, não conta. Se a venda saiu por corretor de fora do time, é venda que estamos perdendo.",
    "'Marcada como venda SDR no CV, sem evidência': o CV tem o tipo 'Vendas SDR', mas o histórico não mostra atendimento recente com repasse. O SDR pode ter",
    "atendido fora do CV; confirmar com o SDR antes de contar ou descartar.",
):
    ws.append([linha])
ws["A1"].font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 100
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 18

COLUNAS = [
    ("Reserva", 9), ("Data da venda", 12), ("Produto", 36), ("Valor (R$)", 13), ("Tipo da venda", 24), ("Corretor", 28), ("Imobiliária", 30), ("Cliente", 30),
    ("Veredito", 44), ("O que falta", 60), ("Atendeu até o fim?", 11), ("Repasse/agendamento", 13), ("Origem online", 11), ("Recente (90d)", 11),
    ("Corretor antes do SDR?", 12), ("Origem paga (campanha)", 12), ("Investimento na praça, 90d (R$)", 16), ("Último atendimento", 14), ("Dias até a venda", 11), ("CV marca SDR?", 11), ("Lead no CV", 11), ("Reserva no CV", 11),
]
sn = lambda b: "sim" if b else "não"


def aba(nome, itens):
    w = wb.create_sheet(nome)
    w.append([c[0] for c in COLUNAS])
    for i, (_, larg) in enumerate(COLUNAS, 1):
        w.column_dimensions[get_column_letter(i)].width = larg
        cel = w.cell(row=1, column=i)
        cel.font, cel.fill, cel.alignment = VERDE, FUNDO, Alignment(wrap_text=True, vertical="center")
    for t in itens:
        w.append([
            t["idreserva"], t["data_venda"], t["produto"], t["valor"], t["canal"], t["corretor"], t["imobiliaria"], t["cliente"],
            t["veredito_label"], "; ".join(t["faltas"]) or "—", sn(t["atendimento"]), sn(t["repasse"]), sn(t["origem"]), sn(t["recente"]),
            sn(t.get("corretor_primeiro")), sn(t.get("origem_paga")), round(t.get("investimento_praca_90d") or 0), t["ultimo_atendimento"].isoformat() if t["ultimo_atendimento"] else None, t["dias_ultimo_atendimento"],
            sn(t["cv_marca_sdr"]), t["link"], t["link_reserva"],
        ])
        r = w.max_row
        w.cell(row=r, column=4).number_format = "#,##0"
        w.cell(row=r, column=17).number_format = "#,##0"
        for col in (21, 22):
            if w.cell(row=r, column=col).value:
                w.cell(row=r, column=col).hyperlink = w.cell(row=r, column=col).value
                w.cell(row=r, column=col).value = "abrir"
                w.cell(row=r, column=col).font = Font(color="0563C1", underline="single")
        for col in range(1, len(COLUNAS) + 1):
            w.cell(row=r, column=col).fill = PatternFill("solid", fgColor=CORES[t["veredito"]])
    w.freeze_panes = "C2"
    w.auto_filter.ref = w.dimensions


aba("Venda SDR", [t for t in todas if t["veredito"] == "sdr"])
aba("Perdidas (saíram por fora)", [t for t in todas if t["veredito"] == "perdida"])
aba("CV marca SDR, sem evidência", [t for t in todas if t["veredito"] == "sdr_cv_sem_evidencia"])
aba("Falta algum requisito", [t for t in todas if t["veredito"] == "parcial"])
aba("Sem rastro", [t for t in todas if t["veredito"] == "sem_rastro"])
aba("Praças remanescentes (fora)", [t for t in todas if t["veredito"] == "remanescente"])
destino = os.path.join("output", "relatorios", f"Vendas_SDR_online_{inicio[:7].replace(chr(45), chr(95))}.xlsx")
wb.save(destino)
print("salvo:", destino, Counter(t["veredito"] for t in todas))
