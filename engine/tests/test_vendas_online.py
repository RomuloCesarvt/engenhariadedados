"""
Testes das vendas online do período (regra da planilha revisada, com a leitura da conversa) e do relatório.
"""

from datetime import date

import pytest
from openpyxl import load_workbook

from engine import bi, identidade, leitura_conversa, origem_rd, relatorio_vendas, vendas_online  # noqa: F401

CP2 = "Avaré - Central Parque II - 13900"
DO_TIME = "Especialista - Patrick Hernani Iaiz"

CONVERSOU = [
    {"data_cad": "2026-07-01 10:00:00", "descricao": "[M.I.A] Lead integrada pela Assistente Virtual"},
    {"data_cad": "2026-07-01 11:00:00", "descricao": "# Relato enviado pela Glória, a assistente virtual da Moura Leite - "
                                                     "Resumo da Conversa: - A lead perguntou sobre o valor da entrada."},
]
VISITA = [{"data_cad": "2026-07-02 10:00:00", "nome": "Visita agendada pela assistente virtual"}]
SEM_RESPOSTA = [
    {"data_cad": "2026-07-01 11:00:00", "descricao": "# Relato enviado pela Glória - A assistente enviou uma mensagem "
                                                     "inicial, mas o lead não respondeu, e não houve interação."},
    {"data_cad": "2026-07-10 11:00:00", "descricao": "Mensagem de sal 2 enviada."},
    {"data_cad": "2026-07-20 11:00:00", "descricao": "Lead não interage de nenhuma maneira."},
]


def lead(idlead, data_cad, midia, nome, email, ult_conv=None, situacao="Em Atendimento", interacao=(), tarefa=(),
         motivo=None):
    return {"idlead": str(idlead), "nome": nome, "email": email, "telefone": "+5514991234567",
            "data_cad": f"{data_cad} 10:00:00", "ultima_data_conversao": f"{ult_conv or data_cad} 10:00:00",
            "midia_principal": midia, "empreendimento": [{"id": 1, "nome": CP2}],
            "link_interesses": f"https://cv/lead/{idlead}", "situacao": {"nome": situacao},
            "interacao": list(interacao), "tarefa": list(tarefa), "motivo_cancelamento": motivo}


def reserva(idreserva, data_venda, tipo_venda=None, leads=(), situacao="Vendida", data_distrato=None,
            corretor="Fulano de Tal", data_reserva="2026-08-18"):
    return {"idreserva": str(idreserva), "situacao": situacao, "vendida": situacao == "Vendida",
            "data": f"{data_reserva} 10:00:00", "data_venda": f"{data_venda} 10:00:00", "data_distrato": data_distrato,
            "empreendimento": CP2, "bloco": "Q1", "unidade": f"L-{idreserva}", "valor_contrato": 200000.0,
            "vgv_tabela": 0.0, "tipo_venda": tipo_venda, "midia": None, "corretor": corretor,
            "imobiliaria": "Imob X", "leads_associados": [str(l) for l in leads], "idprecadastro": None,
            "titular_email": None, "titular_telefone": None}


def bruta(idreserva, nome, email, tipo_venda=None, leads=()):
    return {"titular": {"nome": nome, "email": email, "celular": "+5514991234567"}, "associados_array": [],
            "tipovenda": tipo_venda, "idprecadastro": None, "condicoes": {"valor_contrato": "200000,00"},
            "leads_associados": [{"idlead": str(l)} for l in leads]}


@pytest.fixture
def cenario(monkeypatch, tmp_path):
    """Bases do CV em memória, leitura do CV sem rede e revisões de leitura num arquivo temporário."""
    conteudo = {}
    monkeypatch.setattr(bi.bases, "carregar", lambda chave: conteudo.get(chave))
    monkeypatch.setattr(bi.bases, "registros", lambda chave: (conteudo.get(chave) or {}).get("registros") or [])
    monkeypatch.setattr(vendas_online, "link_reserva", lambda i: f"https://cv/reserva/{i}")
    monkeypatch.setattr(vendas_online, "link_lead", lambda i: f"https://cv/lead/{i}")
    monkeypatch.setattr(relatorio_vendas, "link_lead", lambda i: f"https://cv/lead/{i}")
    monkeypatch.setattr(origem_rd, "por_lead", dict)
    monkeypatch.setattr(leitura_conversa, "ARQUIVO_REVISOES", str(tmp_path / "leituras_revisadas.json"))
    cache = {}
    monkeypatch.setattr(vendas_online, "_completar_cache",
                        lambda nome, pedidos, extrair: {c: cache.get(nome, {}).get(c) for c, _ in pedidos})

    def definir(chave, registros):
        conteudo[chave] = {"registros": registros, "dados_de": "v1", "atualizado_em": f"v{len(conteudo)}",
                           "total": len(registros)}
        bi.limpar_cache()

    definir.cache = cache
    return definir


def ana(**extra):
    return lead(1, "2026-07-01", "Facebook Ads", "Ana Souza Lima", "ana@exemplo.com", **extra)


def ana_bruta(tipo=None, leads=()):
    return {"100": bruta("100", "ANA SOUZA LIMA", "ana@exemplo.com", tipo, leads)}


def _analisar(cenario, leads, reservas, brutas, inicio=date(2026, 8, 1), fim=date(2026, 8, 31)):
    cenario("cv_leads", leads)
    cenario("cv_indice", [])
    cenario("cv_reservas", reservas)
    cenario.cache["reservas"] = brutas
    return vendas_online.analisar(inicio, fim)


def test_venda_sdr_ja_e_online_no_cv(cenario):
    r = _analisar(cenario, [], [reserva(100, "2026-08-20", "Vendas SDR", corretor=DO_TIME)], ana_bruta("Vendas SDR"))
    assert r["vendas"][0]["resultado"] == "online_cv_sdr"
    assert (r["resumo"]["online_cv"], r["resumo"]["casos"]) == (1, 0)


def test_lead_online_vinculado_a_reserva_continua_online_quando_a_conversa_confirma(cenario):
    r = _analisar(cenario, [ana(interacao=CONVERSOU, tarefa=VISITA)],
                  [reserva(100, "2026-08-20", "Venda Time Moura Leite", leads=[1], corretor=DO_TIME)],
                  ana_bruta("Venda Time Moura Leite", leads=[1]))
    venda = r["vendas"][0]
    assert venda["resultado"] == "online_cv_lead" and venda["leitura"]["leitura"] == "online"
    assert (r["resumo"]["online"], r["resumo"]["tirados_do_online"]) == (1, 0)


def test_venda_online_no_cv_sai_do_online_quando_o_lead_nao_respondeu(cenario):
    """Ter mídia online por trás só diz que a venda poderia ser SDR: quem decide é a conversa."""
    r = _analisar(cenario, [ana(interacao=SEM_RESPOSTA)],
                  [reserva(100, "2026-08-20", "Venda Imobiliária", leads=[1])], ana_bruta(leads=[1]))
    venda = r["vendas"][0]
    assert venda["venda_online_cv"] is True
    assert (venda["resultado"], venda["leitura"]["motivo"]) == ("nao_nossa", "sem_resposta")
    assert (r["resumo"]["online"], r["resumo"]["online_cv"]) == (0, 1)
    assert (r["resumo"]["tirados_do_online"], r["resumo"]["vgv_tirados_do_online"]) == (1, 200000.0)


def test_venda_sdr_no_cv_nao_e_relida_mesmo_sem_conversa_registrada(cenario):
    r = _analisar(cenario, [ana(interacao=SEM_RESPOSTA)],
                  [reserva(100, "2026-08-20", "Vendas SDR", leads=[1], corretor=DO_TIME)],
                  ana_bruta("Vendas SDR", leads=[1]))
    assert r["vendas"][0]["resultado"] == "online_cv_sdr" and r["vendas"][0]["leitura_auto"] is None


def test_cliente_que_conversou_e_comprou_com_o_time_e_caso_para_investigar(cenario):
    r = _analisar(cenario, [ana(interacao=CONVERSOU, tarefa=VISITA)],
                  [reserva(100, "2026-08-20", "Venda Time Moura Leite", corretor=DO_TIME)], ana_bruta())
    venda = r["vendas"][0]
    assert venda["resultado"] == "caso" and venda["resultado_label"] == "PARA INVESTIGAR"
    assert "agendou visita" in venda["leitura"]["destaques"]
    assert "comprou com Patrick Hernani Iaiz, do nosso time" in venda["leitura"]["texto"]
    assert (r["resumo"]["casos"], r["resumo"]["vgv_casos"]) == (1, 200000.0)


def test_cliente_que_conversou_e_comprou_com_outro_corretor_e_perda(cenario):
    r = _analisar(cenario, [ana(interacao=CONVERSOU)], [reserva(100, "2026-08-20", "Venda Imobiliária")], ana_bruta())
    assert r["vendas"][0]["resultado"] == "perda"
    assert "mas comprou com Fulano de Tal, da Imob X" in r["vendas"][0]["leitura"]["texto"]
    assert (r["resumo"]["perdas"], r["resumo"]["outro_corretor"]) == (1, 1)


def test_cliente_que_ja_estava_com_outro_corretor_nao_era_venda_nossa(cenario):
    conversa = CONVERSOU + [{"data_cad": "2026-07-05 10:00:00", "descricao": "Já está em cadastro com outro corretor."}]
    r = _analisar(cenario, [ana(interacao=conversa)], [reserva(100, "2026-08-20", "Venda Imobiliária")], ana_bruta())
    venda = r["vendas"][0]
    assert (venda["resultado"], venda["leitura"]["motivo"]) == ("nao_nossa", "ja_com_outro")
    assert "Já está em cadastro com outro corretor." in venda["leitura"]["texto"]


def test_cliente_que_nao_respondeu_nao_era_venda_nossa_mesmo_vendido_pelo_time(cenario):
    r = _analisar(cenario, [ana(interacao=SEM_RESPOSTA)],
                  [reserva(100, "2026-08-20", "Venda Time Moura Leite", corretor=DO_TIME)], ana_bruta())
    assert (r["vendas"][0]["resultado"], r["vendas"][0]["leitura"]["motivo"]) == ("nao_nossa", "sem_resposta")


def test_nao_aceitou_nao_e_sinal_de_conversa():
    lido = leitura_conversa.ler(
        {"data_venda": "2026-08-20", "conv_ref": "2026-07-01", "do_time": False, "corretor": "Fulano"},
        [ana(interacao=[{"data_cad": "2026-07-03 10:00:00",
                         "descricao": "Lead não aceitou os termos quando se cadastrou, então não consegue receber as mensagens."}])])
    assert lido["leitura"] == "nao_nossa"


def test_leitura_revisada_pelo_analista_muda_o_resultado_e_pode_ser_desfeita(cenario):
    r = _analisar(cenario, [ana(interacao=SEM_RESPOSTA)], [reserva(100, "2026-08-20", "Venda Imobiliária")], ana_bruta())
    assert r["vendas"][0]["resultado"] == "nao_nossa"

    leitura_conversa.revisar("100", "perda", "Conversou por telefone com o SDR e agendou visita.")
    revisada = vendas_online.recalcular(r)
    venda = revisada["vendas"][0]
    assert venda["resultado"] == "perda" and venda["leitura"]["revisada"] is True
    assert venda["leitura"]["texto"] == "Conversou por telefone com o SDR e agendou visita."
    assert revisada["resumo"]["leituras_revisadas"] == 1

    assert leitura_conversa.desfazer("100") is True
    assert vendas_online.recalcular(r)["vendas"][0]["resultado"] == "nao_nossa"


def test_leitura_invalida_e_recusada(cenario):
    with pytest.raises(ValueError):
        leitura_conversa.revisar("100", "talvez")


def test_registro_do_corretor_antes_da_conversao_tira_a_venda_dos_casos(cenario):
    comercial = lead(2, "2026-06-01", "Painel Corretor", "Ana Souza Lima", "ana@exemplo.com")
    r = _analisar(cenario, [ana(interacao=CONVERSOU), comercial], [reserva(100, "2026-08-20", "Venda Imobiliária")],
                  ana_bruta())
    venda = r["vendas"][0]
    assert venda["resultado"] == "ja_atendia" and venda["leitura_auto"] is None
    assert "Lead do comercial" in venda["corretor_antes"][0][1]


def test_reserva_colada_na_conversao_sem_sdr_e_corretor_que_ja_atendia(cenario):
    colado = lead(1, "2026-08-16", "Facebook Ads", "Ana Souza Lima", "ana@exemplo.com")
    r = _analisar(cenario, [colado], [reserva(100, "2026-08-20", "Venda Imobiliária", data_reserva="2026-08-18")],
                  ana_bruta())
    assert r["vendas"][0]["resultado"] == "ja_atendia"
    assert "junto com a conversão online" in r["vendas"][0]["corretor_antes"][0][1]


def test_conversao_de_mais_de_90_dias_antes_fica_fora_da_janela(cenario):
    antigo = lead(1, "2026-01-05", "Facebook Ads", "Ana Souza Lima", "ana@exemplo.com")
    r = _analisar(cenario, [antigo], [reserva(100, "2026-08-20", "Venda Time Moura Leite", corretor=DO_TIME)], ana_bruta())
    assert r["vendas"][0]["online_90_dias"] is False
    assert r["vendas"][0]["resultado"] == "fora_janela"


def test_lead_de_dezembro_vale_para_venda_de_janeiro(cenario):
    dezembro = lead(1, "2025-12-10", "Facebook Ads", "Ana Souza Lima", "ana@exemplo.com")
    r = _analisar(cenario, [dezembro], [reserva(100, "2026-01-20", "Venda Time Moura Leite", corretor=DO_TIME,
                                                data_reserva="2026-01-18")],
                  ana_bruta(), inicio=date(2026, 1, 1), fim=date(2026, 1, 31))
    assert r["vendas"][0]["online_90_dias"] is True


def test_venda_distratada_continua_na_analise(cenario):
    distrato = reserva(100, "2026-08-20", "Venda Imobiliária", situacao="Distrato", data_distrato="2026-09-10 10:00:00")
    r = _analisar(cenario, [ana()], [distrato], ana_bruta())
    assert r["resumo"]["vendas"] == 1 and r["resumo"]["distratos"] == 1
    assert r["vendas"][0]["situacao_atual"] == "Distrato"


def test_ligacao_so_por_telefone_com_nome_diferente_e_descartada(cenario):
    carlos = lead(1, "2026-07-01", "Facebook Ads", "Carlos Pereira Alves", "carlos@exemplo.com")
    r = _analisar(cenario, [carlos], [reserva(100, "2026-08-20", "Venda Imobiliária")], ana_bruta())
    assert r["vendas"][0]["leads"] == []
    assert r["vendas"][0]["resultado"] == "sem_conversao"


def test_origem_descoberta_na_rd_torna_o_lead_online(cenario, monkeypatch):
    monkeypatch.setattr(origem_rd, "por_lead", lambda: {
        "1": {"categoria": "Site / landing page (formulário)", "canal": "site", "online": True,
              "detalhe": "1ª conversão na RD: form-lp-avare em 2026-07-01"}})
    sem_origem = lead(1, "2026-07-01", None, "Ana Souza Lima", "ana@exemplo.com")
    r = _analisar(cenario, [sem_origem], [reserva(100, "2026-08-20", "Venda Time Moura Leite", corretor=DO_TIME)],
                  ana_bruta())
    assert r["vendas"][0]["leads"][0]["categoria"] == "Site / landing page (formulário)"
    assert r["vendas"][0]["online_90_dias"] is True


@pytest.mark.parametrize("corretor,esperado", [
    ("Especialista - Patrick Hernani Iaiz", True),
    ("ELISANGELA LOURENCO DE VECCHI", True),
    ("Fulano de Tal", False),
    (None, False),
])
def test_quem_e_do_time(corretor, esperado):
    assert vendas_online.do_time(corretor) is esperado


def test_relatorio_sai_com_as_abas_e_as_vendas_nas_listas_certas(cenario, tmp_path):
    bia = lead(2, "2026-07-02", "Facebook Ads", "Bia Costa Ramos", "bia@exemplo.com", interacao=CONVERSOU)
    brutas = {**ana_bruta(), "101": bruta("101", "BIA COSTA RAMOS", "bia@exemplo.com", "Venda Time Moura Leite")}
    analise = _analisar(cenario, [ana(interacao=CONVERSOU), bia],
                        [reserva(100, "2026-08-20", "Venda Imobiliária"),
                         reserva(101, "2026-08-22", "Venda Time Moura Leite", corretor=DO_TIME)], brutas)
    wb = load_workbook(relatorio_vendas.gerar(analise, caminho=str(tmp_path / "relatorio.xlsx")))
    assert wb.sheetnames == ["Leia primeiro", "Resumo por produto", "Casos para investigar", "Não viraram venda online",
                             "Todas as vendas", "Como foi feito"]
    casos, perdas = wb["Casos para investigar"], wb["Não viraram venda online"]
    assert casos["C5"].value == "BIA COSTA RAMOS" and casos["C5"].hyperlink.target == "https://cv/lead/2"
    assert perdas["C5"].value == "ANA SOUZA LIMA"
    assert perdas["K5"].value.startswith("ATENDEMOS, MAS COMPROU COM OUTRO CORRETOR.")
    cabecalhos = [c.value for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert not [h for h in cabecalhos if "onfiança" in h]
    assert wb["Todas as vendas"].max_row == 3

    tela = relatorio_vendas.para_a_tela(analise)
    assert [v["idreserva"] for v in tela["casos"]] == ["101"]
    assert [(v["idreserva"], v["resultado"]) for v in tela["nao_viraram"]] == [("100", "perda")]
    assert tela["nao_viraram"][0]["leitura"]["leitura_label"] == "Atendemos, mas comprou com outro corretor"


def test_texto_dos_dados_nunca_vira_formula_no_excel(cenario, tmp_path):
    perigoso = lead(1, "2026-07-01", "Facebook Ads", "=SOMA(A1)", "ana@exemplo.com", interacao=CONVERSOU)
    brutas = {"100": bruta("100", "=CMD|' /C calc'!A0", "ana@exemplo.com", "Venda Imobiliária")}
    analise = _analisar(cenario, [perigoso], [reserva(100, "2026-08-20", "Venda Imobiliária")], brutas)
    wb = load_workbook(relatorio_vendas.gerar(analise, caminho=str(tmp_path / "relatorio.xlsx")))
    fora_do_padrao = [c.value for ws in wb for row in ws.iter_rows() for c in row
                      if c.data_type == "f" and not relatorio_vendas.FORMULA_DO_RELATORIO.match(str(c.value))]
    assert fora_do_padrao == []


def test_inicio_da_pre_campanha_e_o_mes_em_que_a_captacao_dispara():
    contagem = {"2025-04": 2, "2025-05": 3, "2025-06": 1, "2025-07": 5, "2025-08": 4, "2025-09": 250, "2025-10": 420}
    assert vendas_online._inicio_pre_campanha(contagem, date(2026, 3, 25)) == date(2025, 9, 1)


def test_produto_sem_salto_de_captacao_nao_tem_pre_campanha():
    contagem = {"2025-06": 4, "2025-07": 5, "2025-08": 6, "2025-09": 5}
    assert vendas_online._inicio_pre_campanha(contagem, date(2025, 10, 1)) is None


@pytest.mark.parametrize("valor,esperado", [
    ("+55 (14) 99123-4567", "1491234567"),
    ("14991234567", "1491234567"),
    ("1491234567", "1491234567"),
    ("991234567", None),
])
def test_chave_de_telefone_ignora_ddi_e_nono_digito(valor, esperado):
    assert vendas_online._chave_telefone(valor) == esperado


@pytest.mark.parametrize("nome,esperado", [
    ("Ana Souza Lima", "mesma pessoa"),
    ("Carla Souza Lima", "familiar (sobrenome em comum)"),
    ("Carlos Pereira Alves", "nome diferente"),
])
def test_compatibilidade_de_nome_com_os_compradores(nome, esperado):
    assert vendas_online._compatibilidade_nome(nome, ["ANA SOUZA LIMA"]) == esperado
