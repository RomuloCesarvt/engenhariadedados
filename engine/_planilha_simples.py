"""Planilha simples: cada venda com uma frase e o link (uso: python engine/_planilha_simples.py 2026-09-01 2026-09-30)."""
import json
import os
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from engine import casos_sdr  # noqa: E402

inicio, fim = sys.argv[1], sys.argv[2]
arq = os.path.join("output", "analises", "periodos", f"{inicio.replace('-', '')}_{fim.replace('-', '')}.json")
todas = casos_sdr.todas(json.load(open(arq, encoding="utf-8"))["analise"])
reais = lambda v: "R$ " + f"{v:,.0f}".replace(",", ".")
CORES = {"sdr": "D5F5E3", "perdida": "FADBD8", "sdr_cv_sem_evidencia": "FFE2B3", "parcial": "FFF8DC", "remanescente": "D6DBDF"}
ABAS = [("Pode ser SDR", "sdr"), ("Pode ser venda perdida", "perdida"), ("Confirmar com o SDR", "sdr_cv_sem_evidencia"),
        ("Provavelmente não é SDR", "parcial"), ("Fora - praça sem anúncio", "remanescente")]
COLUNAS = [("Reserva", 9), ("Data da venda", 13), ("Produto", 38), ("Valor (R$)", 13), ("Quem vendeu", 30), ("Por quê", 95), ("Link do lead", 14), ("Link da reserva", 14), ("Cliente", 30)]
mes = f"{inicio[8:]}/{inicio[5:7]} a {fim[8:]}/{fim[5:7]}/{fim[:4]}"

wb = Workbook()
ws = wb.active
ws.title = "Resumo"
ws.append([f"Vendas que podem ser SDR — {mes}"])
ws.append([])
ws.append(["Vendas do período", len(todas), reais(sum(t["valor"] for t in todas))])
for nome, chave in ABAS:
    g = [t for t in todas if t["veredito"] == chave]
    ws.append([nome, len(g), reais(sum(t["valor"] for t in g))])
ws.append(["Sem sinal de SDR (não listadas)", sum(t["veredito"] == "sem_rastro" for t in todas), ""])
ws.append([])
for l in ("Como ler: cada aba traz as vendas daquele grupo, com a frase do motivo e o link do lead no CV para conferir.",
          "Pode ser SDR: o SDR/assistente atendeu até passar para vendas, o lead veio de origem online, foi recente (até 90 dias) e o nosso time vendeu.",
          "Pode ser venda perdida: mesmo caso, mas a venda saiu por corretor de fora do nosso time.",
          "Confirmar com o SDR: o CV marca como venda SDR, mas o histórico não mostra o atendimento. O SDR pode ter atendido fora do CV."):
    ws.append([l])
ws["A1"].font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 110
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 18

for nome, chave in ABAS:
    w = wb.create_sheet(nome[:31])
    w.append([c[0] for c in COLUNAS])
    for i, (_, larg) in enumerate(COLUNAS, 1):
        w.column_dimensions[get_column_letter(i)].width = larg
        cel = w.cell(row=1, column=i)
        cel.font, cel.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1B4F38")
    for t in [x for x in todas if x["veredito"] == chave]:
        w.append([t["idreserva"], t["data_venda"], t["produto"], t["valor"], t["corretor"], t["explicacao"], t.get("link"), t.get("link_reserva"), t["cliente"]])
        r = w.max_row
        w.cell(row=r, column=4).number_format = "#,##0"
        w.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (7, 8):
            if w.cell(row=r, column=col).value:
                w.cell(row=r, column=col).hyperlink = w.cell(row=r, column=col).value
                w.cell(row=r, column=col).value = "abrir"
                w.cell(row=r, column=col).font = Font(color="0563C1", underline="single")
        for col in range(1, len(COLUNAS) + 1):
            w.cell(row=r, column=col).fill = PatternFill("solid", fgColor=CORES[chave])
    w.freeze_panes = "C2"
destino = os.path.join("output", "relatorios", f"Vendas_possiveis_SDR_{inicio[:7].replace(chr(45), chr(95))}.xlsx")
wb.save(destino)
print("salvo:", destino, dict(Counter(t["veredito"] for t in todas)))
