"""Leads da lista do CV (origem 'Outros'): origem real pelas conversões e sinal de SDR (uso: python engine/_leads_outros.py <csv>)."""
import csv
import json
import os
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from engine import leads_bi  # noqa: E402

L = lambda n: json.load(open(f"output/bases/{n}.json", encoding="utf-8"))["registros"]
linhas = list(csv.DictReader(open(sys.argv[1], encoding="utf-8-sig", errors="replace"), delimiter=";"))
col_ult = [k for k in linhas[0] if k.endswith("ltima Origem")][0]
ids = {int(x["Id"]) for x in linhas}
conv = defaultdict(list)
for c in L("cv_conversoes"):
    if int(c["idlead"]) in ids:
        conv[int(c["idlead"])].append(c)
dw = {int(x["idlead"]): x for x in L("cv_leads_dw") if int(x["idlead"]) in ids}
ONLINE = leads_bi.ORIGENS_ONLINE_BI
PASSOU = {"Em tratativa Corretor", "Lead Vinculado ao Corretor", "Lead Qualificado (SQL)"}


def origens(i):
    vistos = []
    for c in sorted(conv.get(i, []), key=lambda c: str(c.get("referencia_data") or "")):
        o = (c.get("midia_conversao") or "").strip()
        if o and o not in vistos:
            vistos.append(o)
    return vistos


saida = []
for x in linhas:
    i = int(x["Id"])
    ult = x[col_ult]
    og = origens(i)
    online = [o for o in og if leads_bi._chave_origem(o) in ONLINE]
    google = [o for o in og if o.lower().startswith("pesquisa") or "google" in o.lower()]
    sit = (dw.get(i) or {}).get("situacao", "")
    ia = "mia" in x["Tag"].lower() or "morada.ai" in x["E-mail"]
    if ult != "Outros":
        continue
    if sit in PASSOU and ia:
        v, motivo = "Pode ser SDR", f"a IA atendeu o lead e ele hoje está em '{sit}'"
    elif sit == "Em atendimento IA" or sit.startswith("Tentativa"):
        v, motivo = "Em atendimento (ainda sem repasse)", f"está em '{sit}'"
    else:
        v, motivo = "Sem repasse", f"está em '{sit}'"
    if og and not online:
        motivo += "; a origem real é " + ", ".join(og) + (" (a busca do Google, que o BI não conta como online)" if google else " (fora da lista online do BI)")
    elif og:
        motivo += "; origem real: " + ", ".join(og)
    saida.append(dict(id=i, data=x["Data Primeiro Cadastro"], nome=x["Nome"], tel=x["Telefone"].strip("'"), prim=x["Primeira Origem"], real=", ".join(og) or "—",
                      online="sim" if online else "não", sit=sit, ia="sim" if ia else "não", org="sim" if "mia-organico" in x["Tag"] else "não", corretor=x["Corretor"], v=v, motivo=motivo,
                      link=f"https://mouraleite.cvcrm.com.br/gestor/comercial/leads/{i}/administrar"))

CORES = {"Pode ser SDR": "D5F5E3", "Em atendimento (ainda sem repasse)": "FFF8DC", "Sem repasse": "FFFFFF"}
wb = Workbook()
ws = wb.active
ws.title = "Resumo"
c = Counter(s["v"] for s in saida)
ws.append(["Leads da lista com última origem 'Outros'", len(saida)])
for k, n in c.items():
    ws.append([k, n])
ws.append(["Origem real conta como online no BI", sum(s["online"] == "sim" for s in saida)])
ws.append(["Origem real NÃO conta como online no BI", sum(s["online"] == "não" for s in saida)])
ws.column_dimensions["A"].width = 60
COL = [("Id", 9), ("Cadastro", 12), ("Nome", 28), ("Telefone", 16), ("1ª origem", 14), ("Origem real (conversões)", 36), ("Online no BI?", 10), ("Situação no CV", 28), ("IA atendeu?", 10), ("Tag orgânico?", 10), ("Corretor", 24), ("Resultado", 30), ("Por quê", 90), ("Link", 10)]
w = wb.create_sheet("Leads")
w.append([c[0] for c in COL])
for k, (_, larg) in enumerate(COL, 1):
    w.column_dimensions[get_column_letter(k)].width = larg
    w.cell(row=1, column=k).font = Font(bold=True, color="FFFFFF")
    w.cell(row=1, column=k).fill = PatternFill("solid", fgColor="1B4F38")
for s in sorted(saida, key=lambda s: (s["v"] != "Pode ser SDR", s["v"], -s["id"])):
    w.append([s["id"], s["data"], s["nome"], s["tel"], s["prim"], s["real"], s["online"], s["sit"], s["ia"], s["org"], s["corretor"], s["v"], s["motivo"], "abrir"])
    r = w.max_row
    w.cell(row=r, column=14).hyperlink = s["link"]
    w.cell(row=r, column=14).font = Font(color="0563C1", underline="single")
    w.cell(row=r, column=13).alignment = Alignment(wrap_text=True, vertical="top")
    for k in range(1, 14):
        w.cell(row=r, column=k).fill = PatternFill("solid", fgColor=CORES[s["v"]])
w.freeze_panes = "C2"
w.auto_filter.ref = w.dimensions
destino = os.path.join("output", "relatorios", "Leads_origem_outros_Altavista.xlsx")
wb.save(destino)
print(destino, dict(c), Counter(s["online"] for s in saida), Counter(s["ia"] for s in saida))
print(Counter(s["real"] for s in saida if s["online"] == "não").most_common(8))
print(Counter(s["sit"] for s in saida))
