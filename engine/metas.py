"""
Meta de leads por praça e empreendimento, com o que falta passando para o mês seguinte.

A meta vem da planilha "Metas Anuais - comercial e mkt", aba "Memória de cálculo (ano)": um bloco por empreendimento, com "Meta geração de leads
80% / 100% / 120%" mês a mês (junho a dezembro). Usa-se a meta de 100%. Importar:  python -m engine.metas importar "<planilha.xlsx>"

Ponto de partida do saldo: `ajustes_meta_efetiva` guarda a meta efetiva que a agência calculou para um mês (e-mail de 01/09: setembro = Central Parque 334,
Reserva 164, Jardim dos Ipês 151, Novo Horizonte 50). Dali em diante o saldo é calculado pelo sistema com os leads reais.

Regras (pedidas pela gestão):
  - Meta efetiva do mês = meta do mês + o que faltou nos meses anteriores (meta efetiva do mês anterior menos o atingido, se positivo).
  - O que passa do mês anterior é só a falta; sobra não abate a meta do mês seguinte.
  - O saldo começa a contar em `inicio_saldo` (padrão 2026-09: a tabela de conversões do CV só existe a partir de 12/08, então setembro é o primeiro mês completo).

Leads atingidos: os mesmos MQL online do Resumo (regra do Power BI da gestão), por empreendimento.
"""

import calendar
import json
import os
import sys
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Set

from engine import bases, bi

ARQUIVO = os.path.join(bases.BASES_DIR, "metas_leads.json")
INICIO_SALDO_PADRAO = "2026-09"
ABA = "Memória de cálculo (ano)"
MESES_COLUNAS = {"R": "2026-06", "S": "2026-07", "T": "2026-08", "U": "2026-09", "V": "2026-10", "W": "2026-11", "X": "2026-12"}

# Nome do bloco na planilha -> empreendimento do sistema (segundo valor: True quando o vínculo é só uma suposição). "Terras de Itu II" é o Altavista (confirmado pela gestão).
MAPA_PRODUTOS = {
    "RESERVA CENTRAL PARQUE AVARÉ 2": ("Avaré - Reserva Central Parque II - 14000", False),
    "CENTRAL PARQUE AVARÉ 2": ("Avaré - Central Parque II - 13900", False),
    "NOVO HORIZONTE ITATINGA": ("Itatinga - Novo Horizonte - 16103", False),
    "JARDIM DOS IPÊS III": ("Taguaí - Jardim dos Ipês III - 17701", False),
    "RESERVA DOS RAMOS": ("Tatuí - Reserva dos Ramos Residencial - 14704", False),
    "TERRAS DE ITU II": ("Itu - Altavista - 15300", False),
    "GASTÃO DAL FARRA": (None, False),
}


def importar(caminho: str) -> Dict[str, Any]:
    import openpyxl
    from openpyxl.utils import column_index_from_string as col

    ws = openpyxl.load_workbook(caminho, data_only=True)[ABA]
    produtos: Dict[str, Any] = {}
    for r in range(1, ws.max_row + 1):
        nome = ws.cell(r, 1).value
        if not (isinstance(nome, str) and ws.cell(r + 1, 2).value == "Indicador"):
            continue
        metas: Dict[str, Dict[str, int]] = {}
        for k in range(r, r + 22):
            rotulo = str(ws.cell(k, col("Q")).value or "")
            for pct in ("80", "100", "120"):
                if rotulo.startswith("Meta geração de leads") and pct in rotulo:
                    for letra, mes in MESES_COLUNAS.items():
                        v = ws.cell(k, col(letra)).value
                        metas.setdefault(mes, {})[pct] = int(round(v)) if isinstance(v, (int, float)) else 0
        # CPL médio histórico do empreendimento (coluna "2026" do bloco "Projeção lead") e o valor ideal que a planilha usa de referência
        cpl_hist = cpl_ideal = None
        for k in range(r, r + 22):
            if str(ws.cell(k, col("AB")).value or "").strip() == "CPL médio":
                v1, v2 = ws.cell(k, col("AC")).value, ws.cell(k, col("AD")).value
                cpl_hist = float(v1) if isinstance(v1, (int, float)) else None
                cpl_ideal = float(v2) if isinstance(v2, (int, float)) else None
        sistema, assumido = MAPA_PRODUTOS.get(nome.strip(), (None, False))
        produtos[nome.strip()] = {"produto_sistema": sistema, "mapeamento_assumido": assumido, "metas": metas,
                                  "cpl_historico": round(cpl_hist, 2) if cpl_hist else None, "cpl_ideal": round(cpl_ideal, 2) if cpl_ideal else None}
    anterior = lido() or {}
    conteudo = {"fonte": os.path.basename(caminho), "importado_em": datetime.now().isoformat(timespec="seconds"),
                "inicio_saldo": anterior.get("inicio_saldo") or INICIO_SALDO_PADRAO, "meta_usada": "100",
                "ajustes_meta_efetiva": anterior.get("ajustes_meta_efetiva") or {},
                "inicio_saldo_por_produto": anterior.get("inicio_saldo_por_produto") or {"RESERVA DOS RAMOS": "2026-06"}, "produtos": produtos}
    os.makedirs(bases.BASES_DIR, exist_ok=True)
    with open(ARQUIVO, "w", encoding="utf-8") as f:
        json.dump(conteudo, f, ensure_ascii=False, indent=1)
    return conteudo


def lido() -> Optional[Dict[str, Any]]:
    try:
        with open(ARQUIVO, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def _mes_seguinte(mes: str) -> str:
    a, m = int(mes[:4]), int(mes[5:7])
    return f"{a + 1:04d}-01" if m == 12 else f"{a:04d}-{m + 1:02d}"


def _fim_do_mes(ano: int, mes: int) -> date:
    return date(ano, mes, calendar.monthrange(ano, mes)[1])


def _meses(de: str, ate: str) -> List[str]:
    saida, (a, m) = [], (int(de[:4]), int(de[5:7]))
    while f"{a:04d}-{m:02d}" <= ate:
        saida.append(f"{a:04d}-{m:02d}")
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)
    return saida


def calcular(fim: date, emp_f: Optional[Set[str]] = None, praca_f: Optional[Set[str]] = None) -> Optional[Dict[str, Any]]:
    """Meta, saldo e atingido do mês de `fim` (mês até a data), por empreendimento e por praça."""
    from engine.leads_bi import leads_do_bi

    cfg = lido()
    if not cfg:
        return None
    emp_f, praca_f = emp_f or set(), praca_f or set()
    # O atingido só conta até onde o CV foi baixado: com a base atrasada, "até hoje" faria a projeção parecer pior do que é.
    ultima = max((str(c.get("data_cad") or "")[:10] for c in bases.registros("cv_conversoes")), default="")
    if ultima and ultima < fim.isoformat():
        fim = date.fromisoformat(ultima)
    mes_ref = f"{fim.year:04d}-{fim.month:02d}"
    uso = cfg.get("meta_usada") or "100"
    produtos: List[Dict[str, Any]] = []
    for nome_arquivo, p in cfg["produtos"].items():
        sistema = p.get("produto_sistema")
        praca = (bi.praca_do_empreendimento(sistema) if sistema else None) or nome_arquivo.title()
        if (emp_f and sistema not in emp_f) or (praca_f and praca not in praca_f):
            continue
        saldo, historico, atual = 0, [], None
        # Início do saldo por produto (ex.: Reserva dos Ramos, que a agência não atende, conta desde junho); senão o padrão da planilha
        inicio = min((cfg.get("inicio_saldo_por_produto") or {}).get(nome_arquivo) or cfg.get("inicio_saldo") or INICIO_SALDO_PADRAO, mes_ref)
        for mes in _meses(inicio, mes_ref):
            base = int((p["metas"].get(mes) or {}).get(uso, 0))
            efetiva = base + saldo
            informada = ((cfg.get("ajustes_meta_efetiva") or {}).get(mes) or {}).get(nome_arquivo)
            if informada is not None:
                # Meta efetiva já calculada pela agência para o mês (ponto de partida do saldo): vale no lugar da conta
                efetiva, saldo = int(informada), max(0, int(informada) - base)
            ano, m = int(mes[:4]), int(mes[5:7])
            limite = fim if mes == mes_ref else _fim_do_mes(ano, m)
            atingido = len(leads_do_bi(date(ano, m, 1), limite, {sistema}, None)) if sistema else None
            falta = max(0, efetiva - atingido) if atingido is not None else 0
            linha = {"mes": mes, "meta_base": base, "saldo_anterior": saldo, "meta_efetiva": efetiva, "atingido": atingido, "falta": falta}
            if mes == mes_ref:
                atual = linha
            else:
                historico.append(linha)
            saldo = falta
        if atual is None:
            continue
        dias_mes = calendar.monthrange(fim.year, fim.month)[1]
        projecao = round(atual["atingido"] / fim.day * dias_mes) if atual["atingido"] is not None and fim.day else None
        restante_dias = dias_mes - fim.day
        produtos.append({
            "praca": praca, "produto": sistema or nome_arquivo.title(), "arquivo": nome_arquivo, "mapeamento_assumido": bool(p.get("mapeamento_assumido")),
            "sem_dados": sistema is None, **atual,
            "pct": round(atual["atingido"] / atual["meta_efetiva"] * 100, 1) if atual["atingido"] is not None and atual["meta_efetiva"] else None,
            "cpl_historico": p.get("cpl_historico"), "cpl_ideal": p.get("cpl_ideal"),
            "prox_mes": _mes_seguinte(mes_ref), "prox_meta_base": int((p["metas"].get(_mes_seguinte(mes_ref)) or {}).get(uso, 0)),
            "projecao": projecao, "pct_projecao": round(projecao / atual["meta_efetiva"] * 100, 1) if projecao is not None and atual["meta_efetiva"] else None,
            "ritmo_necessario_dia": round(atual["falta"] / restante_dias, 1) if restante_dias > 0 and atual["atingido"] is not None else None,
            "ritmo_atual_dia": round(atual["atingido"] / fim.day, 1) if atual["atingido"] is not None and fim.day else None,
            "inicio_saldo": inicio,
            "historico": historico,
        })

    pracas: Dict[str, Dict[str, Any]] = {}
    for p in produtos:
        g = pracas.setdefault(p["praca"], {"praca": p["praca"], "meta_base": 0, "saldo_anterior": 0, "meta_efetiva": 0, "atingido": 0, "falta": 0, "projecao": 0, "produtos": []})
        if p["sem_dados"]:
            g["meta_base"] += p["meta_base"]; g["meta_efetiva"] += p["meta_efetiva"]; g["sem_dados"] = True
            g["produtos"].append(p)
            continue
        for k in ("meta_base", "saldo_anterior", "meta_efetiva", "atingido", "falta"):
            g[k] += p[k] or 0
        g["projecao"] += p["projecao"] or 0
        g["ritmo_atual_dia"] = round(g.get("ritmo_atual_dia", 0) + (p.get("ritmo_atual_dia") or 0), 1)
        g["ritmo_necessario_dia"] = round(g.get("ritmo_necessario_dia", 0) + (p.get("ritmo_necessario_dia") or 0), 1)
        g["produtos"].append(p)
    lista = []
    for g in pracas.values():
        g["pct"] = round(g["atingido"] / g["meta_efetiva"] * 100, 1) if g["meta_efetiva"] else None
        g["pct_projecao"] = round(g["projecao"] / g["meta_efetiva"] * 100, 1) if g["meta_efetiva"] else None
        lista.append(g)
    lista.sort(key=lambda g: -g["meta_efetiva"])
    return {
        "mes": mes_ref, "ate": fim.isoformat(), "dias_decorridos": fim.day, "dias_no_mes": calendar.monthrange(fim.year, fim.month)[1],
        "inicio_saldo": cfg.get("inicio_saldo") or INICIO_SALDO_PADRAO, "meta_usada": uso, "fonte": cfg.get("fonte"), "pracas": lista,
        # O total soma só praças com meta no mês e com leads no sistema (Itu ainda sem meta e produtos sem leads ficam de fora)
        "total": {k: sum(g[k] for g in lista if g["meta_efetiva"] and not g.get("sem_dados")) for k in ("meta_base", "saldo_anterior", "meta_efetiva", "atingido", "falta", "projecao", "ritmo_atual_dia", "ritmo_necessario_dia")},
        "dados_do_cv_ate": ultima or None,
    }


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "importar":
        r = importar(sys.argv[2])
        print(f"{len(r['produtos'])} produtos importados de {r['fonte']}")
    else:
        print('Uso: python -m engine.metas importar "<planilha.xlsx>"')
